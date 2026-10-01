#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Compêndio: pesquisa + sumário + redação usando GitHub Copilot/Claude."""
import argparse, asyncio, json, math, os, random, re, sys, time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import requests
from bs4 import BeautifulSoup

try:
    from ddgs import DDGS
except Exception:
    DDGS = None

from copilot import CopilotClient
from copilot.session import PermissionHandler
from copilot.session_events import AssistantMessageData

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
OUT = ROOT / "saida"
CACHE = OUT / "_cache"
PARTIAIS = OUT / "_parciais"
ESTADO = OUT / "_estado.json"
CUSTO = OUT / "_custo.json"

PALAVRAS_POR_PAGINA = 400
PALAVRAS_POR_TOPICO = 750
MODELO_PADRAO = "claude-opus-5.5"
PARALELISMO = 6
MAX_FONTES = 12
MAX_TEXTO_FONTE = 12000

EIXOS = [
    ("Apresentação da obra", "ficha técnica, lugar do livro na trajetória do autor, ocasião da escrita, edições e recepção imediata", .04),
    ("Contexto histórico e político", "o Brasil e o mundo no momento da obra; processos que ela precisa explicar", .08),
    ("O debate intelectual", "a conversa acadêmica e pública a que a obra responde; teses que contesta", .08),
    ("Tese central e arquitetura do argumento", "a tese, o problema, as premissas e o encadeamento da demonstração", .08),
    ("Percurso capítulo a capítulo", "cada capítulo: argumento, evidências, função no conjunto e limites", .24),
    ("Conceitos fundamentais", "cada conceito central: definição, origem teórica, aplicação, exemplos e limites", .14),
    ("Método, fontes e evidências", "como o autor pesquisa, que prova mobiliza, força e fragilidade do procedimento", .06),
    ("Diálogos e rupturas", "relação com autores e tradições, um interlocutor por tópico", .08),
    ("Recepção, críticas e controvérsias", "como a obra foi lida, objeções recebidas e respostas possíveis", .08),
    ("Atualidade e aplicações", "o que explica do presente, onde envelheceu, usos em direito e políticas públicas", .06),
    ("Aparato de estudo", "glossário, cronologia, roteiro de leitura, perguntas e leituras complementares", .06),
]

SYSTEM = """Você é um pesquisador acadêmico escrevendo em português brasileiro, com prosa densa, clara e de nível universitário.
REGRAS:
1. Escreva SOMENTE sobre a obra indicada e o tópico solicitado.
2. Não reproduza trechos extensos ou protegidos por direitos autorais. Sintetize e explique com palavras próprias.
3. Não invente páginas, números, citações, capítulos, datas ou argumentos. Se a pesquisa não sustentar um ponto, assinale a incerteza.
4. Diferencie o que o autor sustenta, o que críticos sustentam e o que é interpretação posterior.
5. Não transforme controvérsia acadêmica em fato incontroverso.
6. Não repita material de outros tópicos.
7. Entregue somente o texto final do tópico, sem título, preâmbulo ou despedida.
"""

def log(msg):
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)

def slug(s):
    s = re.sub(r"[^a-z0-9]+", "-", s.lower().encode("ascii", "ignore").decode()).strip("-")
    return s[:90] or "obra"

def load_json(p, default):
    try:
        return json.loads(Path(p).read_text(encoding="utf-8"))
    except Exception:
        return default

def save_json(p, data):
    Path(p).parent.mkdir(parents=True, exist_ok=True)
    tmp = Path(str(p) + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(p)

def load_works():
    p = DATA / "lista_obras.txt"
    return [x.strip() for x in p.read_text(encoding="utf-8").splitlines() if x.strip() and not x.startswith("#")]

def web_search(query, n=5):
    if DDGS is None:
        return []
    try:
        with DDGS() as d:
            return [{"titulo": r.get("title", ""), "url": r.get("href") or r.get("url", ""), "resumo": r.get("body", "")} for r in d.text(query, region="br-pt", max_results=n)]
    except Exception as e:
        log(f"    pesquisa falhou: {str(e)[:80]}")
        return []

def fetch(url):
    try:
        r = requests.get(url, timeout=25, headers={"User-Agent": "Mozilla/5.0 CompendioAcademico/2.0"})
        if r.status_code != 200 or "text/html" not in r.headers.get("content-type", ""):
            return ""
        s = BeautifulSoup(r.text, "html.parser")
        for tag in s(["script", "style", "nav", "footer", "header", "aside", "form"]): tag.decompose()
        t = re.sub(r"\n{2,}", "\n", s.get_text("\n"))
        return re.sub(r"[ \t]{2,}", " ", t).strip()[:MAX_TEXTO_FONTE]
    except Exception:
        return ""

def pesquisar(obra):
    CACHE.mkdir(parents=True, exist_ok=True)
    p = CACHE / f"fontes-{slug(obra)}.json"
    cached = load_json(p, None)
    if cached:
        return cached
    consultas = [
        f'"{obra}" resenha crítica',
        f'"{obra}" resumo capítulos argumento',
        f'"{obra}" conceitos tese análise',
        f'"{obra}" recepção debate críticas',
        f'"{obra}" editora livro',
    ]
    found, seen = [], set()
    for q in consultas:
        for r in web_search(q, 5):
            if r["url"] and r["url"] not in seen:
                seen.add(r["url"]); found.append(r)
    fontes = []
    for r in found[:MAX_FONTES]:
        txt = fetch(r["url"])
        if len(txt) > 500:
            fontes.append({"titulo": r["titulo"], "url": r["url"], "texto": txt})
        elif r["resumo"]:
            fontes.append({"titulo": r["titulo"], "url": r["url"], "texto": r["resumo"]})
    data = {"obra": obra, "fontes": fontes, "em": datetime.now().isoformat()}
    save_json(p, data)
    return data

def contexto(pesquisa, limite=18000):
    partes = []
    for i, f in enumerate(pesquisa.get("fontes", []), 1):
        partes.append(f"[FONTE {i}] {f['titulo']}\nURL: {f['url']}\n{f['texto']}")
    return "\n\n".join(partes)[:limite] or "(Não foram obtidas fontes web suficientes; sinalize incertezas.)"

def make_outline(obra, pesquisa, n):
    quotas = [max(1, round(n * p)) for _, _, p in EIXOS]
    while sum(quotas) > n: quotas[quotas.index(max(quotas))] -= 1
    while sum(quotas) < n: quotas[quotas.index(max(quotas))] += 1
    tops = []
    for (nome, desc, _), q in zip(EIXOS, quotas):
        for i in range(1, q + 1):
            tops.append({"eixo": nome, "titulo": f"{nome} — recorte {i}", "descricao": desc})
    return tops[:n]

async def ask(client, modelo, prompt, timeout=900):
    last = None
    for tentativa in range(5):
        try:
            session = await client.create_session(
                on_permission_request=PermissionHandler.approve_all,
                model=modelo,
            )
            try:
                reply = await session.send_and_wait(prompt, timeout=timeout)
                if reply and isinstance(reply.data, AssistantMessageData):
                    return reply.data.content.strip()
                if reply and hasattr(reply.data, "content"):
                    return str(reply.data.content).strip()
                return ""
            finally:
                try: await session.destroy()
                except Exception: pass
        except Exception as e:
            last = e
            await asyncio.sleep(min(60, 3 * 2 ** tentativa) + random.random())
    log(f"    Claude falhou definitivamente: {str(last)[:100]}")
    return ""

async def gerar_topicos(client, modelo, obra, pesquisa, sumario, parcial, paralelismo):
    dados = load_json(parcial, {"obra": obra, "sumario": sumario, "topicos": {}})
    escritos = dados.setdefault("topicos", {})
    semaforo = asyncio.Semaphore(paralelismo)
    contexto_base = contexto(pesquisa)
    total = len(sumario)

    async def one(i, t):
        key = str(i)
        if escritos.get(key, {}).get("texto", "").strip():
            return
        prompt = f"""OBRA: {obra}
EIXO: {t['eixo']}
TÓPICO: {t['titulo']}
RECORTE: {t['descricao']}
EXTENSÃO: aproximadamente {PALAVRAS_POR_TOPICO} palavras.

MATERIAL DE PESQUISA DISPONÍVEL:
{contexto_base}

TAREFA: produza uma exposição substantiva e específica deste tópico. Não invente detalhes que não estejam sustentados. Quando houver divergência entre fontes, atribua a posição. O resultado será incorporado a um compêndio maior, portanto não faça introdução geral da obra nem repita os demais eixos."""
        async with semaforo:
            text = await ask(client, modelo, SYSTEM + "\n\n" + prompt)
        if text:
            escritos[key] = {"titulo": t["titulo"], "eixo": t["eixo"], "texto": text}
            save_json(parcial, dados)
            log(f"    {obra[:45]} | {i}/{total}")

    for start in range(0, total, paralelismo * 2):
        await asyncio.gather(*(one(i, t) for i, t in list(enumerate(sumario, 1))[start:start + paralelismo * 2]))
    save_json(parcial, dados)
    return dados

def escrever_docx(destino, obra, sumario, escritos, fontes, modelo):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    doc.add_heading(obra, 0)
    p = doc.add_paragraph(f"Compêndio analítico — {modelo} — {datetime.now():%d/%m/%Y}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    atual = None
    for i, t in enumerate(sumario, 1):
        x = escritos.get(str(i))
        if not x: continue
        if x.get("eixo") != atual:
            doc.add_page_break(); doc.add_heading(x.get("eixo", ""), 1); atual = x.get("eixo")
        doc.add_heading(x.get("titulo", t["titulo"]), 2)
        for par in re.split(r"\n+", x.get("texto", "")):
            if par.strip(): doc.add_paragraph(par.strip())
    if fontes:
        doc.add_page_break(); doc.add_heading("Fontes consultadas", 1)
        for f in fontes:
            doc.add_paragraph(f"{f['titulo']} — {f['url']}")
    doc.save(destino)

def contar_palavras(escritos):
    return sum(len(x.get("texto", "").split()) for x in escritos.values())

async def processar(args):
    OUT.mkdir(exist_ok=True); PARTIAIS.mkdir(exist_ok=True); CACHE.mkdir(exist_ok=True)
    obras = load_works()
    if args.pular: obras = obras[args.pular:]
    if args.limite_obras > 0: obras = obras[:args.limite_obras]
    if not obras: raise SystemExit("Nenhuma obra para processar.")

    estado = load_json(ESTADO, {"concluidas": {}})
    client = CopilotClient(github_token=os.getenv("COPILOT_GITHUB_TOKEN") or os.getenv("GITHUB_TOKEN"), use_logged_in_user=False)
    await client.start()
    inicio = time.time(); concluidas = 0
    try:
        for idx, obra in enumerate(obras, args.pular + 1):
            chave = slug(obra)
            destino = OUT / f"{idx:03d}-{chave}.docx"
            if estado.get("concluidas", {}).get(chave) and destino.exists():
                log(f"({idx}) já pronta: {obra}"); continue
            log(f"({idx}) {obra}")
            pesquisa = await asyncio.to_thread(pesquisar, obra)
            n_topicos = math.ceil((args.paginas_por_obra * PALAVRAS_POR_PAGINA) / args.palavras_por_topico)
            sumario = make_outline(obra, pesquisa, n_topicos)
            parcial = PARTIAIS / f"{chave}.json"
            base = load_json(parcial, None)
            if base and base.get("sumario"): sumario = base["sumario"]
            
            dados = await gerar_topicos(client, args.modelo, obra, pesquisa, sumario, parcial, args.paralelismo)
            escritos = dados.get("topicos", {})
            if len(escritos) < len(sumario):
                log(f"  ATENÇÃO: {len(escritos)}/{len(sumario)} tópicos; não marcando como concluída.")
                continue
            escrever_docx(destino, obra, sumario, escritos, pesquisa.get("fontes", []), args.modelo)
            palavras = contar_palavras(escritos)
            estado.setdefault("concluidas", {})[chave] = {"obra": obra, "arquivo": destino.name, "palavras": palavras, "topicos": len(escritos)}
            save_json(ESTADO, estado)
            concluidas += 1
            log(f"  -> {palavras:,} palavras | concluída {concluidas}/{len(obras)}")
    finally:
        await client.stop()
    log(f"FIM. Obras concluídas nesta execução: {concluidas}. Tempo: {(time.time()-inicio)/60:.1f} min")

def main():
    global PALAVRAS_POR_TOPICO
    ap = argparse.ArgumentParser()
    ap.add_argument("--modelo", default=MODELO_PADRAO)
    ap.add_argument("--limite-obras", type=int, default=1)
    ap.add_argument("--pular", type=int, default=0)
    ap.add_argument("--paginas-por-obra", type=int, default=300)
    ap.add_argument("--palavras-por-topico", type=int, default=750)
    ap.add_argument("--paralelismo", type=int, default=PARALELISMO)
    ap.add_argument("--listar", action="store_true")
    args = ap.parse_args()
    PALAVRAS_POR_TOPICO = args.palavras_por_topico
    if args.listar:
        for i, o in enumerate(load_works(), 1): print(f"{i:03d}. {o}")
        return
    asyncio.run(processar(args))

if __name__ == "__main__": main()
